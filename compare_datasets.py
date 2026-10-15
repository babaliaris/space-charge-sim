"""
Compare two sets of curves that were sampled at DIFFERENT x positions
=====================================================================

Typical use: compare curves digitized from a paper (e.g. with WebPlotDigitizer)
against the curves our Python scripts produce.

The problem
-----------
The two files almost never share the same x values, so you cannot subtract
them row by row in Excel. This script lines them up first:

    REFERENCE file : its x values are kept as they are.
                     (normally the digitized paper data)
    COMPARE file   : each of its curves is linearly interpolated at the
                     reference x values.
                     (normally our Python output, which is dense and smooth,
                      so interpolating it adds practically no error)

    difference = compare - reference      (at every reference x)

Reference points that fall outside the x range of the compare curve are
skipped (no extrapolation) and counted in the summary.

Supported file layouts (.csv or .xlsx, detected automatically)
--------------------------------------------------------------
1. WebPlotDigitizer "wide" export: a row of curve names, a row of X,Y labels,
   then one X/Y column pair per curve (curves may have different lengths).
2. Table: one header row, first column = x, every other column = one curve
   (this is what our scripts write, e.g. sheet "Fields" of
   output/wh_fig3_reproduction.xlsx).

Both files must use the same units (e.g. x in cm, y in kV/cm).

How to run
----------
Interactive (asks you to pick files and match curves from numbered lists):

    python compare_datasets.py

Fully from the command line (curves can be given by name or by number):

    python compare_datasets.py --list output/wh_fig3_reproduction.xlsx

    python compare_datasets.py \\
        --reference input/wh_fig3_datasets.csv \\
        --compare   output/wh_fig3_reproduction.xlsx \\
        --pair "Esp=E_sp space charge [kV/cm]" \\
        --pair "Ei,rod=E_I,rod rod image [kV/cm]" \\
        --pair "Esp+Ei,rod=E_sp + E_I,rod + E_I,plane [kV/cm]"

Output (in output/ unless --output is given)
--------------------------------------------
    sheet "Summary"   one row per matched pair: bias, mean absolute error,
                      RMS error, maximum error and where it happens, ...
    one sheet per pair: x, reference y, compare y (interpolated), difference
    sheet "Sources"   which files/sheets/curves were used and when
"""

import argparse
import csv
import re
import sys
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

import numpy as np
from openpyxl import load_workbook

from excel_export import save_to_excel


# =============================================================================
# SETTINGS
# =============================================================================
PROJECT_FOLDER = Path(__file__).resolve().parent
SEARCH_FOLDERS = [PROJECT_FOLDER / "input", PROJECT_FOLDER / "output"]  # offered in interactive mode
OUTPUT_FOLDER  = PROJECT_FOLDER / "output"
FILE_TYPES     = (".csv", ".xlsx")


@dataclass
class Curve:
    name: str
    x: np.ndarray
    y: np.ndarray


# =============================================================================
# 1. READING FILES  ->  list of curves
# =============================================================================
def read_csv_rows(path):
    """Read a .csv file as a list of rows (lists of text cells).

    Handles both ',' files and the ';' + decimal-comma files that Excel writes
    with Greek/European regional settings.
    """
    text = Path(path).read_text(encoding="utf-8-sig")
    lines = text.splitlines()

    # Decide the column separator from the data lines (header names may contain commas)
    sample = "\n".join(lines[2:12])
    delimiter = next((d for d in ("\t", ";") if d in sample), ",")

    rows = list(csv.reader(lines, delimiter=delimiter))
    if delimiter != ",":
        # "1,25" -> "1.25" (only safe when ',' is not the column separator)
        rows = [[cell.replace(",", ".") if _looks_like_decimal_comma(cell) else cell
                 for cell in row] for row in rows]
    return rows


def _looks_like_decimal_comma(cell):
    return bool(re.fullmatch(r"\s*[-+]?\d*,\d+([eE][-+]?\d+)?\s*", cell))


def read_xlsx_rows(path, sheet_name):
    """Read one sheet of an .xlsx file as a list of rows."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    rows = [list(row) for row in workbook[sheet_name].iter_rows(values_only=True)]
    workbook.close()
    return rows


def xlsx_sheet_names(path):
    workbook = load_workbook(path, read_only=True)
    names = workbook.sheetnames
    workbook.close()
    return names


def to_number(cell):
    """Return the cell as a float, or None if it is empty / not a number."""
    if cell is None or isinstance(cell, bool):
        return None
    if isinstance(cell, (int, float)):
        return float(cell)
    try:
        return float(str(cell).strip())
    except ValueError:
        return None


def is_blank(cell):
    return cell is None or str(cell).strip() == ""


def parse_curves(rows):
    """Turn raw rows into curves, detecting which of the two layouts is used."""
    rows = [row for row in rows if any(not is_blank(c) for c in row)]   # drop empty lines
    if len(rows) < 2:
        return []
    if _is_webplotdigitizer_layout(rows):
        return _parse_webplotdigitizer(rows)
    return _parse_table(rows)


def _is_webplotdigitizer_layout(rows):
    labels = [str(c).strip().upper() for c in rows[1] if not is_blank(c)]
    return len(labels) >= 2 and labels == ["X", "Y"] * (len(labels) // 2)


def _parse_webplotdigitizer(rows):
    n_curves = len([c for c in rows[1] if not is_blank(c)]) // 2

    # Curve names: WebPlotDigitizer writes "name,," per curve. A name that itself
    # contains a comma (e.g. "Ei,rod") gets split, so we glue together the pieces
    # between empty cells.
    names, piece = [], []
    for cell in rows[0]:
        if is_blank(cell):
            if piece:
                names.append(",".join(piece))
                piece = []
        else:
            piece.append(str(cell).strip())
    if piece:
        names.append(",".join(piece))
    if len(names) != n_curves:
        names = [f"Curve {i + 1}" for i in range(n_curves)]

    curves = []
    for i, name in enumerate(names):
        points = []
        for row in rows[2:]:
            x = to_number(row[2 * i]) if len(row) > 2 * i else None
            y = to_number(row[2 * i + 1]) if len(row) > 2 * i + 1 else None
            if x is not None and y is not None:
                points.append((x, y))
        if points:
            curves.append(_make_curve(name, points))
    return curves


def _parse_table(rows):
    header = rows[0]
    curves = []
    for col in range(1, len(header)):
        points = []
        for row in rows[1:]:
            x = to_number(row[0]) if row else None
            y = to_number(row[col]) if len(row) > col else None
            if x is not None and y is not None:
                points.append((x, y))
        if points:
            name = str(header[col]).strip() if not is_blank(header[col]) else f"Column {col + 1}"
            curves.append(_make_curve(name, points))
    return curves


def _make_curve(name, points):
    x, y = np.array(points, dtype=float).T
    return Curve(name, x, y)


def load_curves(path, sheet=None):
    """Load all curves of a file. For .xlsx, use `sheet` or the first sheet that has curves.

    Returns (curves, sheet name used or None).
    """
    path = Path(path)
    if path.suffix.lower() == ".csv":
        return parse_curves(read_csv_rows(path)), None

    if path.suffix.lower() == ".xlsx":
        sheets = [sheet] if sheet else xlsx_sheet_names(path)
        for name in sheets:
            curves = parse_curves(read_xlsx_rows(path, name))
            if curves:
                return curves, name
        return [], sheet

    raise ValueError(f"Unsupported file type '{path.suffix}' (use .csv or .xlsx).")


# =============================================================================
# 2. COMPARING ONE PAIR OF CURVES
# =============================================================================
def compare_curves(reference, compare):
    """Interpolate `compare` at the x values of `reference` and compute differences."""
    # np.interp needs increasing x: sort the compare curve
    order = np.argsort(compare.x)
    cmp_x, cmp_y = compare.x[order], compare.y[order]

    # Only reference points inside the compare curve's x range (no extrapolation)
    inside = (reference.x >= cmp_x[0]) & (reference.x <= cmp_x[-1])
    x = reference.x[inside]
    ref_y = reference.y[inside]
    cmp_y_at_x = np.interp(x, cmp_x, cmp_y)

    difference = cmp_y_at_x - ref_y
    abs_difference = np.abs(difference)
    peak_reference = np.max(np.abs(ref_y)) if len(ref_y) else np.nan

    table = {
        "x": x,
        f"reference: {reference.name}": ref_y,
        f"compare: {compare.name} (interpolated)": cmp_y_at_x,
        "difference (compare - reference)": difference,
        "|difference|": abs_difference,
        "difference / peak |reference| [%]": 100.0 * difference / peak_reference,
    }

    if len(x):
        worst = int(np.argmax(abs_difference))
        summary = {
            "points compared": len(x),
            "points skipped (outside compare x range)": int(np.sum(~inside)),
            "x from": float(x.min()),
            "x to": float(x.max()),
            "mean difference (bias)": float(np.mean(difference)),
            "mean |difference| (MAE)": float(np.mean(abs_difference)),
            "RMS difference": float(np.sqrt(np.mean(difference ** 2))),
            "max |difference|": float(abs_difference[worst]),
            "x of max |difference|": float(x[worst]),
            "peak |reference|": float(peak_reference),
            "MAE / peak |reference| [%]": float(100.0 * np.mean(abs_difference) / peak_reference),
            "RMS / peak |reference| [%]": float(100.0 * np.sqrt(np.mean(difference ** 2)) / peak_reference),
        }
    else:
        summary = {"points compared": 0,
                   "points skipped (outside compare x range)": int(np.sum(~inside))}
    return table, summary


# =============================================================================
# 3. CHOOSING FILES AND CURVES (interactive prompts / command line)
# =============================================================================
def ask(prompt):
    try:
        return input(prompt).strip()
    except EOFError:
        sys.exit("\nNo input - stopped.")


def ask_number(prompt, maximum, allow_empty=False):
    while True:
        answer = ask(prompt)
        if allow_empty and answer == "":
            return None
        if answer.isdigit() and 1 <= int(answer) <= maximum:
            return int(answer)
        print(f"  Please type a number from 1 to {maximum}" + (" (or press Enter to skip)." if allow_empty else "."))


def available_files():
    files = []
    for folder in SEARCH_FOLDERS:
        if folder.is_dir():
            files += sorted(p for p in folder.iterdir()
                            if p.suffix.lower() in FILE_TYPES and not p.name.startswith("~$"))
    return files


def pick_file_interactively(role, files):
    print(f"\nChoose the {role} file:")
    for i, path in enumerate(files, 1):
        print(f"  [{i}] {path.relative_to(PROJECT_FOLDER)}")
    path = files[ask_number("Number: ", len(files)) - 1]

    sheet = None
    if path.suffix.lower() == ".xlsx":
        sheets = xlsx_sheet_names(path)
        if len(sheets) > 1:
            print(f"  '{path.name}' has several sheets:")
            for i, name in enumerate(sheets, 1):
                print(f"    [{i}] {name}")
            sheet = sheets[ask_number("  Sheet number: ", len(sheets)) - 1]
    return path, sheet


def print_curves(title, curves):
    print(f"\n{title}")
    for i, curve in enumerate(curves, 1):
        print(f"  [{i}] {curve.name}   ({len(curve.x)} points, x {curve.x.min():g} to {curve.x.max():g})")


def pick_pairs_interactively(reference_curves, compare_curves_list):
    print_curves("Curves in the COMPARE file:", compare_curves_list)
    print("\nFor each REFERENCE curve, type the number of the COMPARE curve to match it with"
          "\n(press Enter to skip that curve).")
    pairs = []
    for ref in reference_curves:
        number = ask_number(f"  '{ref.name}' -> ", len(compare_curves_list), allow_empty=True)
        if number is not None:
            pairs.append((ref, compare_curves_list[number - 1]))
    return pairs


def find_curve(curves, key, file_role):
    """Find a curve by number (1, 2, ...) or by name (exact, or a unique start of the name)."""
    key = key.strip()
    if key.isdigit() and 1 <= int(key) <= len(curves):
        return curves[int(key) - 1]
    exact = [c for c in curves if c.name.lower() == key.lower()]
    if exact:
        return exact[0]
    starts = [c for c in curves if c.name.lower().startswith(key.lower())]
    if len(starts) == 1:
        return starts[0]
    names = ", ".join(f"'{c.name}'" for c in curves)
    problem = "matches several curves" if starts else "matches no curve"
    sys.exit(f"--pair: '{key}' {problem} in the {file_role} file. Available: {names}")


def pairs_from_arguments(pair_arguments, reference_curves, compare_curves_list):
    pairs = []
    for text in pair_arguments:
        if "=" not in text:
            sys.exit(f"--pair '{text}' must look like REFERENCE=COMPARE")
        ref_key, cmp_key = text.split("=", 1)
        pairs.append((find_curve(reference_curves, ref_key, "reference"),
                      find_curve(compare_curves_list, cmp_key, "compare")))
    return pairs


def safe_sheet_name(text, used):
    """Excel sheet names: max 31 characters, none of []:*?/\\ , and unique."""
    name = re.sub(r"[\[\]:*?/\\]", "", text).strip()[:31] or "Sheet"
    base, counter = name, 2
    while name.lower() in used:
        suffix = f" ({counter})"
        name = base[:31 - len(suffix)] + suffix
        counter += 1
    used.add(name.lower())
    return name


# =============================================================================
# 4. MAIN
# =============================================================================
def main():
    parser = argparse.ArgumentParser(
        description="Compare curves sampled at different x values (interpolates the COMPARE "
                    "file onto the x values of the REFERENCE file). Run without arguments "
                    "for interactive mode.")
    parser.add_argument("--reference", help="file whose x values are kept (e.g. digitized paper data)")
    parser.add_argument("--compare", help="file whose curves are interpolated (e.g. our Python output)")
    parser.add_argument("--reference-sheet", help="sheet of the reference .xlsx (default: first sheet with curves)")
    parser.add_argument("--compare-sheet", help="sheet of the compare .xlsx (default: first sheet with curves)")
    parser.add_argument("--pair", action="append", default=[], metavar="REF=CMP",
                        help="match a reference curve with a compare curve, by name or number; repeat for more pairs")
    parser.add_argument("--output", help="output .xlsx path (default: output/comparison_<reference>_vs_<compare>.xlsx)")
    parser.add_argument("--list", metavar="FILE", help="only list the curves found in FILE and stop")
    parser.add_argument("--sheet", help="sheet to use with --list")
    args = parser.parse_args()

    # --- Just list the curves of a file ---------------------------------------
    if args.list:
        curves, sheet = load_curves(args.list, args.sheet)
        print_curves(f"Curves in {args.list}" + (f" (sheet '{sheet}')" if sheet else "") + ":", curves)
        return

    # --- Step 1: choose the two files ------------------------------------------
    files = available_files()
    if args.reference:
        reference_path, reference_sheet = Path(args.reference), args.reference_sheet
    else:
        if not files:
            sys.exit("No .csv/.xlsx files found in input/ or output/.")
        reference_path, reference_sheet = pick_file_interactively("REFERENCE (x values kept)", files)
    if args.compare:
        compare_path, compare_sheet = Path(args.compare), args.compare_sheet
    else:
        compare_path, compare_sheet = pick_file_interactively("COMPARE (interpolated onto the reference x)", files)

    reference_curves, reference_sheet = load_curves(reference_path, reference_sheet)
    compare_curves_list, compare_sheet = load_curves(compare_path, compare_sheet)
    if not reference_curves:
        sys.exit(f"No curves found in {reference_path}.")
    if not compare_curves_list:
        sys.exit(f"No curves found in {compare_path}.")

    # --- Step 2: match the curves ------------------------------------------------
    if args.pair:
        pairs = pairs_from_arguments(args.pair, reference_curves, compare_curves_list)
    else:
        print_curves("Curves in the REFERENCE file:", reference_curves)
        pairs = pick_pairs_interactively(reference_curves, compare_curves_list)
    if not pairs:
        sys.exit("No curves were matched - nothing to compare.")

    # --- Step 3: compare each pair -----------------------------------------------
    print("\nResults:")
    sheets = {}
    summary_rows = []
    used_names = {"summary", "sources"}
    for number, (ref, cmp) in enumerate(pairs, 1):
        table, summary = compare_curves(ref, cmp)
        sheets[safe_sheet_name(f"{number}. {ref.name}", used_names)] = table
        summary_rows.append({"pair": number, "reference curve": ref.name, "compare curve": cmp.name, **summary})

        if summary["points compared"]:
            print(f"  [{number}] {ref.name}  vs  {cmp.name}\n"
                  f"      {summary['points compared']} points compared, "
                  f"{summary['points skipped (outside compare x range)']} skipped | "
                  f"bias {summary['mean difference (bias)']:+.3f}, MAE {summary['mean |difference| (MAE)']:.3f}, "
                  f"RMS {summary['RMS difference']:.3f}, max {summary['max |difference|']:.3f} "
                  f"at x = {summary['x of max |difference|']:g}")
        else:
            print(f"  [{number}] {ref.name}  vs  {cmp.name}: no overlapping x range, nothing compared")
        if len(cmp.x) < len(ref.x):
            print("      note: the COMPARE curve has fewer points than the REFERENCE curve; "
                  "swapping the two files usually gives a more accurate interpolation.")

    # Summary sheet: one row per pair (missing values left blank)
    columns = list(dict.fromkeys(key for row in summary_rows for key in row))
    summary_sheet = {column: [row.get(column) for row in summary_rows] for column in columns}

    sources_sheet = {
        "Item": ["Reference file", "Reference sheet", "Compare file", "Compare sheet",
                 "Method", "Difference", "Units", "Generated on"],
        "Value": [str(reference_path), reference_sheet or "-", str(compare_path), compare_sheet or "-",
                  "compare curves linearly interpolated at the reference x values; no extrapolation",
                  "compare - reference",
                  "same as the y values of the input files",
                  datetime.now().strftime("%Y-%m-%d %H:%M")],
    }

    # --- Step 4: save -----------------------------------------------------------
    output_path = Path(args.output) if args.output else \
        OUTPUT_FOLDER / f"comparison_{reference_path.stem}_vs_{compare_path.stem}.xlsx"
    save_to_excel(output_path, {"Summary": summary_sheet, **sheets, "Sources": sources_sheet})
    print(f"\nSaved to {output_path}")


if __name__ == "__main__":
    main()
