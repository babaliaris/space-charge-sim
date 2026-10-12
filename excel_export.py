"""
Helper for saving script results to an Excel (.xlsx) workbook.

Every sheet is given as a dictionary of columns:

    {"z from tip apex [cm]": [...], "E_total [kV/cm]": [...]}

The keys become the bold header row and the values the column data, so each
sheet opens in Excel as a clean table that is ready to plot.
"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font


def save_to_excel(file_path, sheets):
    """
    Write one workbook with one sheet per entry of `sheets`.

    file_path : where to save the .xlsx file (parent folders are created).
    sheets    : {sheet name: {column header: list of values}}.
                All columns of a sheet must have the same length.
    """
    file_path = Path(file_path)
    file_path.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    workbook.remove(workbook.active)  # drop the empty default sheet

    for sheet_name, columns in sheets.items():
        sheet = workbook.create_sheet(title=sheet_name)
        headers = list(columns.keys())
        values = [list(column) for column in columns.values()]

        lengths = {len(column) for column in values}
        if len(lengths) > 1:
            raise ValueError(f"Sheet '{sheet_name}': all columns must have the same length.")

        # Header row in bold
        sheet.append(headers)
        for cell in sheet[1]:
            cell.font = Font(bold=True)

        # Data rows (numpy numbers are turned into plain Python numbers for Excel)
        for row in zip(*values):
            sheet.append([value.item() if hasattr(value, "item") else value for value in row])

        # Keep the header visible while scrolling, and make columns wide enough to read
        sheet.freeze_panes = "A2"
        for index, header in enumerate(headers):
            longest = max([len(str(header))] + [len(f"{v:.6g}" if isinstance(v, float) else str(v))
                                                for v in values[index]])
            sheet.column_dimensions[sheet.cell(row=1, column=index + 1).column_letter].width = longest + 2

    try:
        workbook.save(file_path)
    except PermissionError:
        raise PermissionError(f"Could not write '{file_path}'. Is it open in Excel? "
                              "Close it and run the script again.") from None
